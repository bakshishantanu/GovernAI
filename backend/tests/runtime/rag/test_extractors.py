import io

import pytest

from app.runtime.rag.extractors import (
    ExtractionFailed,
    UnsupportedFileType,
    detect_format,
    extract_pdf,
)


def test_detect_format_accepts_a_pdf_by_extension():
    assert detect_format("annual-report.pdf", "application/pdf") == "pdf"


def test_extension_wins_over_a_wrong_mime_type():
    """Browsers routinely send application/octet-stream for a PDF; rejecting
    on that alone would refuse valid uploads."""
    assert detect_format("report.PDF", "application/octet-stream") == "pdf"


def test_a_wrong_mime_type_cannot_change_a_known_extension():
    """The other direction of the same distrust: a .docx labelled
    application/pdf must be treated as a Word document, not handed to the PDF
    extractor, which would fail minutes later inside the background job."""
    assert detect_format("notes.docx", "application/pdf") == "docx"


def test_mime_type_is_used_when_the_filename_has_no_extension():
    assert detect_format("report", "application/pdf") == "pdf"


@pytest.mark.parametrize(
    "filename,expected",
    [("paper.docx", "docx"), ("deck.pptx", "pptx"), ("report.PDF", "pdf")],
)
def test_every_supported_office_format_is_recognised(filename, expected):
    assert detect_format(filename, None) == expected


@pytest.mark.parametrize("filename", ["photo.gif", "data.csv", "notes.doc", "archive.zip"])
def test_unsupported_types_are_refused_by_name(filename):
    """.doc is the pre-2007 Word format, which is a different container
    entirely. An upload must be refused up front rather than accepted and
    failed later."""
    with pytest.raises(UnsupportedFileType):
        detect_format(filename, None)


@pytest.mark.parametrize("filename", ["scan.jpg", "photo.PNG", "shot.webp"])
def test_images_are_accepted_for_ocr(filename):
    assert detect_format(filename, None) == "image"


def test_the_refusal_message_names_what_is_supported():
    with pytest.raises(UnsupportedFileType) as exc:
        detect_format("data.csv", None)
    message = str(exc.value)
    assert ".pdf" in message and ".docx" in message and ".pptx" in message


async def test_a_file_that_is_not_a_pdf_reports_extraction_failure():
    with pytest.raises(ExtractionFailed):
        await extract_pdf(b"this is plainly not a pdf")


async def test_blank_pages_are_skipped_rather_than_stored_as_empty_chunks():
    """With no OCR provider, a page with no extractable text is skipped
    rather than stored as an empty chunk."""
    pytest.importorskip("pypdf")
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)

    assert await extract_pdf(buffer.getvalue()) == []


# --- Word ---------------------------------------------------------------
#
# Headings are the locator for a .docx, because the format stores no page
# numbers: pages are produced by whatever renders the file. Getting heading
# detection wrong collapses a whole paper into one unciteable section, which
# is exactly what happened on the first real document tried against it.


def _docx_bytes(build) -> bytes:
    pytest.importorskip("docx")
    import docx

    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _bold(document, text):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    return paragraph


def test_docx_sections_are_keyed_on_real_heading_styles():
    from app.runtime.rag.extractors import extract_docx

    def build(d):
        d.add_heading("Introduction", level=1)
        d.add_paragraph("Monkeypox is a growing health problem.")
        d.add_heading("Methodology", level=1)
        d.add_paragraph("We used episodic training.")

    sections = extract_docx(_docx_bytes(build))

    assert [s.locator for s in sections] == ["Introduction", "Methodology"]
    assert all(s.page_number is None for s in sections)


def test_a_short_bold_paragraph_counts_as_a_heading():
    """Real papers rarely use Word's Heading styles. In the document this was
    built against, all 81 paragraphs were style 'Normal' and every heading was
    bold alone."""
    from app.runtime.rag.extractors import extract_docx

    def build(d):
        _bold(d, "1. Introduction")
        d.add_paragraph("The very fact that monkeypox is an emergency.")
        _bold(d, "2. Related Work")
        d.add_paragraph("Ensemble networks have demonstrated potential.")

    sections = extract_docx(_docx_bytes(build))

    assert [s.locator for s in sections] == ["1. Introduction", "2. Related Work"]


def test_a_long_bold_paragraph_is_body_text_not_a_heading():
    """An abstract is often entirely bold. Promoting it to a section heading
    would put a 90-word paragraph inside every citation."""
    from app.runtime.rag.extractors import extract_docx

    abstract = "Abstract " + " ".join(f"word{i}" for i in range(60))

    def build(d):
        _bold(d, "1. Introduction")
        _bold(d, abstract)

    sections = extract_docx(_docx_bytes(build))

    assert [s.locator for s in sections] == ["1. Introduction"]
    assert "word59" in sections[0].text


def test_text_before_the_first_heading_is_kept():
    from app.runtime.rag.extractors import extract_docx

    def build(d):
        d.add_paragraph("Author list and affiliations.")
        d.add_heading("Introduction", level=1)
        d.add_paragraph("Body.")

    sections = extract_docx(_docx_bytes(build))

    assert sections[0].locator is None
    assert "Author list" in sections[0].text


def test_docx_tables_are_extracted_and_cited_separately():
    """Results tables carry much of a paper's substance and are not in
    `paragraphs` at all, so they would be silently lost."""
    from app.runtime.rag.extractors import extract_docx

    def build(d):
        d.add_heading("Results", level=1)
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Model"
        table.cell(0, 1).text = "Accuracy"
        table.cell(1, 0).text = "ResNet50"
        table.cell(1, 1).text = "0.91"

    sections = extract_docx(_docx_bytes(build))

    tables = [s for s in sections if s.locator == "table 1"]
    assert len(tables) == 1
    assert "ResNet50 | 0.91" in tables[0].text


def test_a_file_that_is_not_a_docx_reports_extraction_failure():
    from app.runtime.rag.extractors import extract_docx

    with pytest.raises(ExtractionFailed):
        extract_docx(b"definitely not a word document")


# --- PowerPoint ---------------------------------------------------------


def _pptx_bytes(build) -> bytes:
    pytest.importorskip("pptx")
    from pptx import Presentation

    presentation = Presentation()
    build(presentation)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_slides_are_cited_as_slides_not_pages():
    """A deck's unit is the slide. Citing 'p.7' for slide 7 would make a
    reader translate, and is simply the wrong noun."""
    from app.runtime.rag.extractors import extract_pptx

    def build(p):
        for title in ("What is Normalization?", "Third Normal Form"):
            slide = p.slides.add_slide(p.slide_layouts[5])
            slide.shapes.title.text = title

    slides = extract_pptx(_pptx_bytes(build))

    assert [s.locator for s in slides] == ["slide 1", "slide 2"]
    assert [s.page_number for s in slides] == [1, 2]


def test_empty_slides_are_skipped_without_shifting_later_slide_numbers():
    """Slide numbering must match what the presenter sees, so a blank slide
    cannot renumber everything after it."""
    from app.runtime.rag.extractors import extract_pptx

    def build(p):
        p.slides.add_slide(p.slide_layouts[6])  # blank
        slide = p.slides.add_slide(p.slide_layouts[5])
        slide.shapes.title.text = "BCNF"

    slides = extract_pptx(_pptx_bytes(build))

    assert [s.locator for s in slides] == ["slide 2"]


def test_speaker_notes_are_included():
    """Notes routinely carry the explanation the slide only gestures at."""
    from app.runtime.rag.extractors import extract_pptx

    def build(p):
        slide = p.slides.add_slide(p.slide_layouts[5])
        slide.shapes.title.text = "Summary"
        slide.notes_slide.notes_text_frame.text = "Mention the BCNF decomposition."

    slides = extract_pptx(_pptx_bytes(build))

    assert "Mention the BCNF decomposition." in slides[0].text


def test_a_file_that_is_not_a_pptx_reports_extraction_failure():
    from app.runtime.rag.extractors import extract_pptx

    with pytest.raises(ExtractionFailed):
        extract_pptx(b"definitely not a presentation")
